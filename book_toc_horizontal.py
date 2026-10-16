# -*- coding: utf-8 -*-
"""横排目录样式：整节横排、从左到右。

用法：python book_toc_horizontal.py [目录文件.md]
生成与源文件同名的 "_横排排版.docx"。

源文件格式（第一行是「目录」两个字，其余每行一条：标题 + 页码）：
    目录
    回忆父亲 01
    在清华的夏天 45
    ...

页码目前是占位假数据，后续根据书籍实际章节替换成真实页码即可。
"""
import io, os, sys
if (getattr(sys.stdout, "encoding", "") or "").lower() != "utf-8":
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")

from docx import Document
from docx.shared import Pt, Mm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml.ns import qn

# ==================== 样式参数（想改样式只改这里） ====================
SRC = sys.argv[1] if len(sys.argv) > 1 else r"D:\纸质文集\目录.md"
NAME = os.path.splitext(os.path.basename(SRC))[0] + "_横排排版"
OUT_DIR = os.path.dirname(os.path.abspath(SRC))

PAGE_W, PAGE_H = 162, 229        # 纸张大小（毫米）：宽 × 高
M_TOP, M_BOT = 32, 40            # 上、下边距（毫米）
M_LEFT, M_RIGHT = 28, 33         # 左、右边距（毫米）
F_TOC = "方正清刻本悦宋简体"     # 目录用到的中文字体
TITLE_SIZE = 20                  # 「目录」两个字的大小（磅）
ENTRY_SIZE = 12                  # 每个条目文字的大小（磅）
TOP_LINES = 0                    # 「目录」从页面顶部往下空几行（1 行 = ENTRY_SIZE 磅）
TITLE_AFTER = 18                 # 「目录」和第一条之间空多少（磅）
ENTRY_LEADER = WD_TAB_LEADER.SPACES  # 标题和页码之间的填充线：DOTS=点线 / SPACES=空白 / HYPHENS=短横线
ENTRY_SPACING = 1.6              # 条目的行距倍数（1.0 最紧凑，越大越松散）
# ======================================================================


def R(p, text, size, bold=False):
    """给段落 p 加一段文字 run，设置好字体和大小。"""
    r = p.add_run(text)
    r.font.name = "Times New Roman"   # 西文（数字/字母）用这个字体
    r.font.size = Pt(size)
    r.font.bold = bold
    r._element.rPr.rFonts.set(qn("w:eastAsia"), F_TOC)  # 中文用目录字体
    return r


def build():
    lines = [l.strip() for l in open(SRC, encoding="utf-8").read().split("\n") if l.strip()]
    title = lines[0]
    entries = []
    for l in lines[1:]:
        parts = l.split()
        entries.append(("".join(parts[:-1]), parts[-1]))

    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(ENTRY_SIZE)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), F_TOC)

    sec = doc.sections[0]
    sec.page_width = Mm(PAGE_W); sec.page_height = Mm(PAGE_H)
    sec.top_margin = Mm(M_TOP); sec.bottom_margin = Mm(M_BOT)
    sec.left_margin = Mm(M_LEFT); sec.right_margin = Mm(M_RIGHT)
    for g in sec._sectPr.findall(qn("w:docGrid")):
        sec._sectPr.remove(g)   # 去掉文档网格，避免行距被强制拉大

    # 「目录」标题：从页顶往下空 TOP_LINES 行，水平居中
    tp = doc.add_paragraph()
    tp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tp.paragraph_format.space_before = Pt(TOP_LINES * ENTRY_SIZE)
    tp.paragraph_format.space_after = Pt(TITLE_AFTER)
    tp.paragraph_format.line_spacing = 1.0
    R(tp, title, TITLE_SIZE, bold=True)

    # 每条目录：标题左对齐，页码右对齐（中间用填充线连接）
    text_w_mm = PAGE_W - M_LEFT - M_RIGHT   # 文字区域宽度 = 页码右对齐的位置
    for t, n in entries:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = ENTRY_SPACING
        p.paragraph_format.tab_stops.add_tab_stop(
            Mm(text_w_mm), WD_TAB_ALIGNMENT.RIGHT, ENTRY_LEADER)
        R(p, t, ENTRY_SIZE)
        p.add_run("\t")
        R(p, str(int(n)), ENTRY_SIZE)

    out = os.path.join(OUT_DIR, f"{NAME}.docx")
    doc.save(out)
    print(f"已生成 {out}")
    print(f"条目 {len(entries)} 条")


if __name__ == "__main__":
    build()
