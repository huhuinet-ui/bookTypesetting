# -*- coding: utf-8 -*-
"""把「前言 + 目录 + 正文 + 后记」合成一本书（一个 Word 文档）。

用法：python book_full.py ["要合成的书.md" ["书名"]]
不带参数时用下方 SRC/BOOK_NAME 默认值；输出的 docx 自动生成在源文件同目录、同名 .docx。

源文件约定写法见 book_layout.py 顶部注释：
  # 自序 / # 前言 / # 序言 / # 代序 / # 写在前面   —— 前言页（排在目录之前）
  # 后记 / # 跋         —— 后记页（排在正文之后）
  ### 章节 + 简介：      —— 章节页（章节名 + 下划线 + 文本框引言）
  #### 标题：+ 作者：    —— 文章（无作者则省略标题下划线和作者名）
  图片：+ 图片名：       —— 插图 + 图注

成书顺序：前言页 → 目录页 → 正文 → 后记页（前言/目录不显示页码，正文从第 1 页开始；
目录里「前言/自序」只列标题、不写页码）。
目录页码自动计算：脚本会先单独排一遍正文并转 PDF，量出每章/每篇的真实起始页，再写进目录。
改样式：分别去改 book_layout.py 和 book_toc_horizontal.py 顶部参数区。
"""
import io, os, re, sys, subprocess, tempfile
if (getattr(sys.stdout, "encoding", "") or "").lower() != "utf-8":
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")

from docx.shared import Pt, Mm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement, parse_xml

import fitz
import book_layout as book
import book_toc_horizontal as toc

# ==================== 输入与输出（命令行传参，带默认值） ====================
SRC = sys.argv[1] if len(sys.argv) > 1 else r"D:\纸质文集\你总是有选择的.md"   # 正文源文件
BOOK_NAME = sys.argv[2] if len(sys.argv) > 2 else "你总是有选择的"              # 页眉上显示的书名
OUT = os.path.join(os.path.dirname(os.path.abspath(SRC)),
                   os.path.splitext(os.path.basename(SRC))[0] + ".docx")        # 输出到源文件同目录、同名
TOC_TITLE = "目录"                          # 目录页大标题

# ---- 目录里章节/文章行的样式 ----
CHAPTER_TOC_SIZE = 13        # 目录里章节行的字号（磅）
CHAPTER_TOC_BOLD = True      # 目录里章节行是否加粗
ARTICLE_INDENT_PT = 24       # 目录里文章行相对章节行缩进多少（磅）
# ==========================================================


def add_toc_line(doc, text, page, size, bold=False, indent_pt=0, pad=None):
    """目录里加一行：标题左对齐（可缩进/加粗），页码右对齐（按 pad 补零）。"""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = toc.ENTRY_SPACING
    if indent_pt:
        p.paragraph_format.left_indent = Pt(indent_pt)
    text_w_mm = book.PAGE_W - book.M_LEFT - book.M_RIGHT
    p.paragraph_format.tab_stops.add_tab_stop(
        Mm(text_w_mm), WD_TAB_ALIGNMENT.RIGHT, toc.ENTRY_LEADER)
    toc.R(p, text, size, bold=bold)
    p.add_run("\t")
    num = "" if not page else (str(page).zfill(pad) if pad else str(page))
    toc.R(p, num, size)


def build_toc(doc, chapters, page_map, pad=None, preface_title=None):
    """在文档里排目录页（横排）。preface_title：前言/自序，只列标题、不写页码。"""
    tp = doc.add_paragraph()
    tp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tp.paragraph_format.space_before = Pt(toc.TOP_LINES * toc.ENTRY_SIZE)
    tp.paragraph_format.space_after = Pt(toc.TITLE_AFTER)
    tp.paragraph_format.line_spacing = 1.0
    toc.R(tp, TOC_TITLE, toc.TITLE_SIZE, bold=True)

    if preface_title:
        add_toc_line(doc, preface_title, "", CHAPTER_TOC_SIZE, bold=CHAPTER_TOC_BOLD)

    for ch in chapters:
        if ch["title"] is not None:
            add_toc_line(doc, ch["title"], page_map.get(ch["title"], ""),
                         CHAPTER_TOC_SIZE, bold=CHAPTER_TOC_BOLD, pad=pad)
        for a in ch["articles"]:
            add_toc_line(doc, a["title"], page_map.get(a["title"], ""),
                         toc.ENTRY_SIZE, indent_pt=ARTICLE_INDENT_PT, pad=pad)


def build_body(doc, chapters, first_sec, pad=None, page_map=None):
    """在 doc 里排正文（第一个正文节用 first_sec，正文页码从 1 开始）。
    篇章页单独成节、置空页眉（暗码：不显示页码）；章内文章合为一节、带奇偶页眉。
    page_map：每篇标题 → 正文起始页（用于判断该页是奇数页/偶数页，好让标题色块镜像）。"""
    set_page_number_start(first_sec._sectPr, 1)   # 只有正文第一节重置页码，其余节连续编号

    def _recto(title):
        # 正文页码 1 起：奇数页=右页(recto)→色块在右/标题右对齐；偶数页=左页(verso)→色块在左/标题左对齐
        if not page_map:
            return True
        pg = page_map.get(title, 1)
        return pg % 2 == 1

    for ci, ch in enumerate(chapters):
        if ch["title"] is not None:
            if ci == 0:
                page_sec = first_sec
            else:
                page_sec = doc.add_section(WD_SECTION.ODD_PAGE)   # 篇章页从奇数页(右页)开始，与参考书一致
                remove_page_number_start(page_sec._sectPr)
            book.clear_header(page_sec)                      # 暗码：篇章页不显示页码
            book.add_page_background(page_sec, str(book.BOOK_COLOR))   # 篇章页整页底色（用书籍主色调）
            book.add_chapter_page(doc, ch["title"], ch.get("intro"))
            art_sec = doc.add_section(WD_SECTION.NEW_PAGE)
            remove_page_number_start(art_sec._sectPr)
            book.setup_header(art_sec, chapter=ch["title"], pad=pad)   # 文章节奇偶页眉
            for i, a in enumerate(ch["articles"]):
                book.build_article(doc, a, i == 0, recto=_recto(a["title"]))
        else:
            book.setup_header(first_sec, pad=pad)
            for i, a in enumerate(ch["articles"]):
                book.build_article(doc, a, i == 0, recto=_recto(a["title"]))


def add_front_matter(doc, title, text):
    """前言/后记页：单独成节、空页眉（暗码，不显示页码）。"""
    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    remove_page_number_start(sec._sectPr)
    book.clear_header(sec)
    book.add_front_page(doc, title, text)
    return sec


def remove_page_number_start(sectPr):
    """移除节里的页码重置（w:pgNumType），让它接着上一节连续编号。"""
    pg = sectPr.find(qn("w:pgNumType"))
    if pg is not None:
        sectPr.remove(pg)


def set_page_number_start(sectPr, start=1):
    """让这一节的页码从 start 开始编号（放在 w:cols 之前，符合 Word 元素顺序）。"""
    pg = OxmlElement("w:pgNumType")
    pg.set(qn("w:start"), str(start))
    cols = sectPr.find(qn("w:cols"))
    if cols is not None:
        cols.addprevious(pg)
    else:
        sectPr.append(pg)


def measure_pages(chapters):
    """单独排一遍正文、转 PDF，量出每章/每篇的起始页（正文页码，从 1 起）。
    返回 (page_map, total_pages)：total_pages 用于算页码补零位数。"""
    tmp_docx = os.path.join(tempfile.gettempdir(), "_bookfull_body.docx")
    tmp_pdf = os.path.join(tempfile.gettempdir(), "_bookfull_body.pdf")

    doc, sec = book.new_doc()
    build_body(doc, chapters, sec)
    doc.save(tmp_docx)

    ps = ("$ErrorActionPreference = 'Stop'; "
          "try { "
          "$w = New-Object -ComObject Word.Application; $w.Visible = $false; "
          f"$d = $w.Documents.Open('{tmp_docx}'); $d.Fields.Update(); $d.Repaginate(); "
          f"$d.SaveAs([ref]'{tmp_pdf}', [ref]17); $d.Close($false); $w.Quit(); "
          "} catch { Write-Output ('WORD_ERROR: ' + $_.Exception.Message); exit 1 }")
    r = subprocess.run(["powershell.exe", "-NoProfile", "-Command", ps], capture_output=True)
    if r.returncode != 0:
        raw = (r.stdout or b"") + b"\n" + (r.stderr or b"")
        try:
            msg = raw.decode("gbk")
        except UnicodeDecodeError:
            msg = raw.decode("utf-8", "replace")
        raise RuntimeError("Word 转 PDF 失败，详情：\n" + msg)

    d = fitz.open(tmp_pdf)
    total_pages = len(d)
    full = [re.sub(r"\s+", "", d[i].get_text()) for i in range(total_pages)]
    titles = []
    for ch in chapters:
        if ch["title"] is not None:
            titles.append(ch["title"])
        titles.extend(a["title"] for a in ch["articles"])
    page_map = {}
    for t in titles:
        tn = re.sub(r"\s+", "", t)
        for i, txt in enumerate(full):
            if tn in txt:
                page_map[t] = i + 1
                break
    d.close()
    os.remove(tmp_docx)
    os.remove(tmp_pdf)
    return page_map, total_pages


def main():
    book.SRC = SRC
    book.BOOK_NAME = BOOK_NAME

    data = book.parse_book()
    chapters = data["chapters"]
    page_map, total_pages = measure_pages(chapters)
    pad = len(str(total_pages)) if total_pages > 0 else 0   # 页码补零位数（87页→2位，123页→3位）

    doc, sec1 = book.new_doc()      # 第一节 = 前言页（无前言时就是目录页）

    # 前言/自序页排在最前（目录之前），暗码（不显示页码）
    if data["preface"] is not None:
        book.clear_header(sec1)
        book.add_front_page(doc, data["preface_title"], data["preface"])
        toc_sec = doc.add_section(WD_SECTION.NEW_PAGE)   # 目录页另起一页
        remove_page_number_start(toc_sec._sectPr)
        book.clear_header(toc_sec)
    else:
        toc_sec = sec1

    build_toc(doc, chapters, page_map, pad=pad, preface_title=data["preface_title"])

    sec2 = doc.add_section(WD_SECTION.ODD_PAGE)   # 正文第一节：从奇数页(右页)开始，保证正文页码奇偶=物理页奇偶，色块/标题镜像才和页眉一致
    build_body(doc, chapters, sec2, pad=pad, page_map=page_map)

    if data["afterword"] is not None:  # 最后一篇文章后插后记页
        add_front_matter(doc, data["afterword_title"], data["afterword"])

    if book.CROP_MARKS:                # 每页四角加裁切角线
        for sec in doc.sections:
            book.add_crop_marks(sec)

    doc.save(OUT)
    n_ch = sum(1 for c in chapters if c["title"] is not None)
    n_art = sum(len(c["articles"]) for c in chapters)
    print(f"已生成 {OUT}")
    print(f"章节 {n_ch} 个；文章 {n_art} 篇；前言={'有' if data['preface'] else '无'}；"
          f"后记={'有' if data['afterword'] else '无'}；正文共 {total_pages} 页；"
          f"页码补零 {pad} 位；目录页码已自动计算")


if __name__ == "__main__":
    main()
